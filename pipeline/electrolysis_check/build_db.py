"""Build the offline retrieval DB for the electrolysis experiment."""
from __future__ import annotations

import json
import time
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"
FEATURES_NPZ = DATA_DIR / "features.npz"
METADATA_JSON = DATA_DIR / "metadata.json"

DATASET_DIR = HERE.parents[2] / "化小北-普化部分" / "电解"
EXCEL_PATH = DATASET_DIR / "详细说明.xlsx"
IMAGE_DIRS = {
    "correct": DATASET_DIR / "正确",
    "wrong": DATASET_DIR / "错误",
}
SHEET_BY_LABEL = {
    "correct": "正确",
    "wrong": "错误",
}

MAX_SIDE = 1600
AKAZE = cv2.AKAZE_create()


def _read_excel_descriptions(excel_path: Path) -> dict[str, dict[str, str]]:
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    descriptions: dict[str, dict[str, str]] = {"correct": {}, "wrong": {}}

    for label, sheet_name in SHEET_BY_LABEL.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing sheet: {sheet_name}")
        sheet = workbook[sheet_name]
        header = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        try:
            name_idx = header.index("图片名")
            desc_idx = header.index("说明")
        except ValueError as exc:
            raise ValueError(f"{sheet_name} sheet must contain 图片名 and 说明 columns.") from exc

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[name_idx] is None:
                continue
            image_name = str(row[name_idx]).strip()
            description = "" if row[desc_idx] is None else str(row[desc_idx]).strip()
            if image_name:
                descriptions[label][image_name] = description

    return descriptions


def _preprocess(image: np.ndarray) -> np.ndarray:
    h, w = image.shape[:2]
    if max(h, w) <= MAX_SIDE:
        working = image.copy()
    else:
        scale = MAX_SIDE / max(h, w)
        working = cv2.resize(
            image,
            (max(1, round(w * scale)), max(1, round(h * scale))),
            interpolation=cv2.INTER_AREA,
        )

    gray = cv2.cvtColor(working, cv2.COLOR_BGR2GRAY)
    clahe = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8))
    return clahe.apply(gray)


def _extract_features(image_path: Path) -> tuple[np.ndarray, np.ndarray | None]:
    data = np.fromfile(str(image_path), dtype=np.uint8)
    image = cv2.imdecode(data, cv2.IMREAD_COLOR)
    if image is None:
        return np.zeros((0, 2), dtype=np.float32), None

    gray = _preprocess(image)
    keypoints, descriptors = AKAZE.detectAndCompute(gray, None)
    if not keypoints or descriptors is None:
        return np.zeros((0, 2), dtype=np.float32), None

    points = np.array([kp.pt for kp in keypoints], dtype=np.float32)
    return points, descriptors.astype(np.uint8)


def _iter_images(image_dir: Path) -> list[Path]:
    exts = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
    return sorted(path for path in image_dir.iterdir() if path.is_file() and path.suffix.lower() in exts)


def build_db() -> dict[str, Any]:
    if not DATASET_DIR.exists():
        raise FileNotFoundError(f"Dataset not found: {DATASET_DIR}")
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel not found: {EXCEL_PATH}")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    descriptions = _read_excel_descriptions(EXCEL_PATH)

    samples: list[dict[str, Any]] = []
    arrays: dict[str, np.ndarray] = {}
    started = time.perf_counter()

    for label, image_dir in IMAGE_DIRS.items():
        if not image_dir.exists():
            raise FileNotFoundError(f"Image dir not found: {image_dir}")

        for image_path in _iter_images(image_dir):
            points, descriptors = _extract_features(image_path)
            if descriptors is None or len(points) < 4:
                print(f"SKIP {label}/{image_path.name}: too few keypoints ({len(points)})")
                continue

            description = descriptions[label].get(image_path.stem, "")
            if not description:
                print(f"WARN no Excel description for {label}/{image_path.name}")

            sample_id = f"{label}_{image_path.stem}"
            pts_key = f"{sample_id}_pts"
            desc_key = f"{sample_id}_desc"
            arrays[pts_key] = points.astype(np.float32)
            arrays[desc_key] = descriptors

            samples.append({
                "sample_id": sample_id,
                "label": label,
                "image_file": image_path.name,
                "image_path": str(image_path),
                "description": description,
                "pts_key": pts_key,
                "desc_key": desc_key,
                "keypoint_count": int(len(points)),
            })

    np.savez_compressed(FEATURES_NPZ, **arrays)
    metadata = {
        "version": 1,
        "feature": "AKAZE",
        "max_side": MAX_SIDE,
        "dataset_dir": str(DATASET_DIR),
        "sample_count": len(samples),
        "label_counts": {
            "correct": sum(1 for sample in samples if sample["label"] == "correct"),
            "wrong": sum(1 for sample in samples if sample["label"] == "wrong"),
        },
        "total_keypoints": sum(sample["keypoint_count"] for sample in samples),
        "samples": samples,
    }
    METADATA_JSON.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    elapsed = (time.perf_counter() - started) * 1000
    print(
        "Done. "
        f"{metadata['label_counts']['correct']} correct, "
        f"{metadata['label_counts']['wrong']} wrong, "
        f"{len(samples)} samples saved in {elapsed:.0f} ms."
    )
    return metadata


if __name__ == "__main__":
    build_db()
